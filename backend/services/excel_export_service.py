import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

def export_customers_to_excel(rows):
    """
    Xuất danh sách khách hàng (từ gateway_sms_logs) ra file Excel.
    Rows là kết quả query từ SQLAlchemy có các cột:
    - phone_number
    - detected_provider
    - total_sent
    - latest_status
    - latest_message
    - last_sent_at
    - latest_device_id
    - latest_routing_strategy
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    
    headers = [
        "STT", "Số điện thoại", "Tên khách hàng", "Email", "Nhà mạng",
        "Số lần đã gửi", "Trạng thái gửi gần nhất", "Nội dung tin nhắn gần nhất",
        "Ngày gửi gần nhất", "Thiết bị gửi gần nhất", "Chiến lược định tuyến", "Ghi chú"
    ]
    
    # Write headers
    ws.append(headers)
    
    # Style headers
    header_font = Font(bold=True)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    
    # Freeze top row
    ws.freeze_panes = "A2"
    
    # Write data
    for i, row in enumerate(rows, start=1):
        # Format phone number to preserve leading zero
        phone_cell = ws.cell(row=i+1, column=2)
        phone_cell.value = str(row.phone_number) if row.phone_number else ""
        phone_cell.number_format = '@' # Text format
        
        ws.cell(row=i+1, column=1).value = i # STT
        ws.cell(row=i+1, column=3).value = "" # Tên khách hàng (chưa có)
        ws.cell(row=i+1, column=4).value = "" # Email (chưa có)
        ws.cell(row=i+1, column=5).value = row.detected_provider or ""
        ws.cell(row=i+1, column=6).value = row.total_sent or 0
        ws.cell(row=i+1, column=7).value = row.latest_status or ""
        ws.cell(row=i+1, column=8).value = row.latest_message or ""
        
        last_sent_at_str = row.last_sent_at.strftime("%Y-%m-%d %H:%M:%S") if row.last_sent_at else ""
        ws.cell(row=i+1, column=9).value = last_sent_at_str
        
        ws.cell(row=i+1, column=10).value = row.latest_device_id or ""
        ws.cell(row=i+1, column=11).value = row.latest_routing_strategy or ""
        ws.cell(row=i+1, column=12).value = "" # Ghi chú
    
    # Add summary row at the end
    total_row = len(rows)
    ws.append([])
    ws.append(["Tổng số khách hàng:", total_row])
    ws.cell(row=total_row+3, column=1).font = Font(bold=True)
    ws.cell(row=total_row+3, column=2).font = Font(bold=True)
    
    # Auto width
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                val = str(cell.value)
                if len(val) > max_length:
                    max_length = len(val)
            except:
                pass
        adjusted_width = (max_length + 2)
        if adjusted_width > 50:
            adjusted_width = 50 # max width
        ws.column_dimensions[column].width = adjusted_width
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_sms_template_excel():
    """
    Tạo file Excel mẫu để import danh sách gửi SMS.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "SMS_Template"
    
    headers = ["phone_number", "customer_name", "message", "note"]
    ws.append(headers)
    
    header_font = Font(bold=True)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
    
    ws.freeze_panes = "A2"
    
    # example row 1
    cell1 = ws.cell(row=2, column=1)
    cell1.value = "0987654321"
    cell1.number_format = '@'
    ws.cell(row=2, column=2).value = "Nguyen Van A"
    ws.cell(row=2, column=3).value = "Chuc mung sinh nhat anh A!"
    ws.cell(row=2, column=4).value = "Khach VIP"

    # example row 2
    cell2 = ws.cell(row=3, column=1)
    cell2.value = "0912345678"
    cell2.number_format = '@'
    ws.cell(row=3, column=2).value = "Tran Thi B"
    ws.cell(row=3, column=3).value = "Thong bao uu dai giam 50%"
    ws.cell(row=3, column=4).value = ""
    
    # Auto width
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                val = str(cell.value)
                if len(val) > max_length:
                    max_length = len(val)
            except:
                pass
        adjusted_width = (max_length + 2)
        if adjusted_width < 15:
            adjusted_width = 15
        ws.column_dimensions[column].width = adjusted_width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
